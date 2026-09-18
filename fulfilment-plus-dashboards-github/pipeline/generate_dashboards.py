"""
Generate two standalone HTML dashboards from kpi_dashboard.xlsx:

  - customer_dashboard.html  - pick a customer, see their KPI scorecards,
                                trend charts and action-plan notes.
  - business_dashboard.html  - whole-business view: overall KPI performance
                                across all customers, per-customer comparison,
                                and trends over time.

Both files are self-contained (Chart.js loaded from CDN, data embedded as
JSON) - just double-click to open in a browser. Re-run this script after
each monthly extraction (extract_kpi_data.py) to refresh them.

Usage:
    python generate_dashboards.py
"""

import os
import json
import datetime
import openpyxl

WORKBOOK           = "kpi_dashboard.xlsx"
MTD_FILE           = "mtd_kpis.json"
DIFOT_CARRIER_FILE = "difot_carriers.json"
PICK_STATS_FILE    = "pick_stats.json"


def load_mtd(path=MTD_FILE):
    if not os.path.exists(path):
        return None
    with open(path) as f:
        return json.load(f)


def load_difot_carriers(path=DIFOT_CARRIER_FILE):
    if not os.path.exists(path):
        return None
    with open(path) as f:
        return json.load(f)


def load_pick_stats(path=PICK_STATS_FILE):
    """Loads pick/pack time stats produced by extract_pick_stats.py (pulled
    from the separate time-tracking app, keyed by customerName - which does
    not always exactly match the KPI dashboard's customer names, hence the
    fuzzy matching done client-side in the dashboard JS)."""
    if not os.path.exists(path):
        return None
    with open(path) as f:
        return json.load(f)


def load_data(path=WORKBOOK):
    wb = openpyxl.load_workbook(path, data_only=True)

    ws = wb["KPI_Data"]
    records = []
    for row in ws.iter_rows(min_row=4, max_row=ws.max_row, values_only=True):
        month, customer, kpi, carrier, target, num, den = row[0:7]
        notes = row[9] if len(row) > 9 else None
        if month is None or customer is None:
            continue
        actual = (num / den) if den else None
        met = (actual is not None) and (actual >= target)
        records.append({
            "month": month.strftime("%Y-%m") if hasattr(month, "strftime") else str(month),
            "customer": customer,
            "kpi": kpi,
            "carrier": carrier,
            "target": target,
            "num": num,
            "den": den,
            "actual": actual,
            "met": met,
            "notes": notes,
        })

    ws2 = wb["Customers"]
    customers = []
    for r in ws2.iter_rows(min_row=4, max_row=ws2.max_row, values_only=True):
        if r[0]:
            customers.append(r[0])

    ws3 = wb["KPI Schedule"]
    kpi_schedule = []
    for r in ws3.iter_rows(min_row=5, max_row=ws3.max_row, values_only=True):
        if not r[0]:
            continue
        kpi_schedule.append({
            "kpi": r[0],
            "target": r[1],
            "frequency": r[2],
            "measure": r[3],
            "notes": r[4],
        })

    return {
        "records": records,
        "customers": customers,
        "kpi_schedule": kpi_schedule,
        "mtd": load_mtd(),
        "difot_carriers": load_difot_carriers(),
        "pick_stats": load_pick_stats(),
    }


PAGE_SHELL = """<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{title}</title>
<script src="https://cdnjs.cloudflare.com/ajax/libs/Chart.js/4.4.1/chart.umd.min.js"></script>
<style>
  :root {{
    --bg: #f4f6f8;
    --card: #ffffff;
    --text: #1f2933;
    --muted: #6b7280;
    --border: #e2e8f0;
    --accent: #1d4ed8;
    --good: #16a34a;
    --good-bg: #dcfce7;
    --bad: #dc2626;
    --bad-bg: #fee2e2;
    --neutral-bg: #f1f5f9;
  }}
  * {{ box-sizing: border-box; }}
  body {{
    margin: 0;
    font-family: -apple-system, Segoe UI, Roboto, Helvetica, Arial, sans-serif;
    background: var(--bg);
    color: var(--text);
  }}
  header {{
    background: #ffffff;
    border-bottom: 1px solid var(--border);
    padding: 20px 28px;
    display: flex;
    flex-wrap: wrap;
    align-items: center;
    justify-content: space-between;
    gap: 12px;
  }}
  header h1 {{ margin: 0; font-size: 1.4rem; }}
  header .meta {{ color: var(--muted); font-size: 0.85rem; }}
  nav.tabs {{ display: flex; gap: 8px; }}
  nav.tabs a {{
    padding: 6px 14px;
    border-radius: 6px;
    text-decoration: none;
    font-size: 0.85rem;
    color: var(--muted);
    border: 1px solid var(--border);
    background: #fff;
  }}
  nav.tabs a.active {{ color: var(--accent); border-color: var(--accent); font-weight: 600; }}
  main {{ padding: 24px 28px 60px; max-width: 1200px; margin: 0 auto; }}
  select {{
    font-size: 1rem;
    padding: 8px 12px;
    border-radius: 6px;
    border: 1px solid var(--border);
    background: #fff;
    color: var(--text);
  }}
  .controls {{ margin-bottom: 20px; display: flex; align-items: center; gap: 12px; flex-wrap: wrap; }}
  .grid {{
    display: grid;
    grid-template-columns: repeat(auto-fit, minmax(260px, 1fr));
    gap: 16px;
    margin-bottom: 28px;
  }}
  .card {{
    background: var(--card);
    border: 1px solid var(--border);
    border-radius: 10px;
    padding: 16px 18px;
  }}
  .card h3 {{ margin: 0 0 4px; font-size: 0.95rem; }}
  .card .freq {{ color: var(--muted); font-size: 0.75rem; margin-bottom: 10px; }}
  .card .figure {{ font-size: 2rem; font-weight: 700; line-height: 1; }}
  .card .target {{ color: var(--muted); font-size: 0.8rem; margin-top: 4px; }}
  .badge {{
    display: inline-block;
    font-size: 0.75rem;
    font-weight: 600;
    padding: 2px 8px;
    border-radius: 999px;
    margin-top: 10px;
  }}
  .badge.good {{ background: var(--good-bg); color: var(--good); }}
  .badge.bad {{ background: var(--bad-bg); color: var(--bad); }}
  .badge.neutral {{ background: var(--neutral-bg); color: var(--muted); }}
  .card.met {{ border-left: 4px solid var(--good); }}
  .card.unmet {{ border-left: 4px solid var(--bad); }}
  .card.nodata {{ border-left: 4px solid var(--border); opacity: 0.7; }}
  .card .note {{ font-size: 0.8rem; color: var(--bad); margin-top: 10px; }}
  .mtd {{
    margin-top: 12px;
    padding-top: 10px;
    border-top: 1px dashed var(--border);
  }}
  .mtd .mtd-label {{ font-size: 0.7rem; color: var(--muted); text-transform: uppercase; letter-spacing: 0.04em; margin-bottom: 2px; }}
  .mtd .mtd-figure {{ font-size: 1.2rem; font-weight: 700; }}
  .mtd .mtd-frac {{ font-size: 0.75rem; font-weight: 400; color: var(--muted); }}
  .mtd .mtd-na {{ font-size: 0.8rem; color: var(--muted); }}
  .ring-card {{ display: flex; flex-direction: column; align-items: center; text-align: center; }}
  .ring-card h3 {{ margin: 10px 0 2px; font-size: 0.95rem; }}
  .ring-card .ring-target-label {{ color: var(--muted); font-size: 0.78rem; margin-bottom: 2px; }}
  .ring-card .ring-last-label {{ color: var(--muted); font-size: 0.78rem; }}
  .ring-track {{ stroke: var(--border); stroke-width: 7; }}
  .ring-track-dashed {{ stroke-dasharray: 4 6; }}
  .ring-progress {{ stroke-width: 7; stroke-linecap: round; transition: stroke-dashoffset 0.4s ease; }}
  .ring-progress.ring-good {{ stroke: var(--good); }}
  .ring-progress.ring-bad {{ stroke: var(--bad); }}
  .ring-pct {{ font-size: 19px; font-weight: 700; fill: var(--text); }}
  .ring-frac {{ font-size: 11px; fill: var(--muted); }}
  .ring-empty {{ font-size: 11px; fill: var(--muted); }}
  .chart-section {{ margin-bottom: 32px; }}
  .chart-section h2 {{ font-size: 1.05rem; margin: 0 0 12px; }}
  .chart-card {{
    background: var(--card);
    border: 1px solid var(--border);
    border-radius: 10px;
    padding: 16px 18px;
    margin-bottom: 16px;
  }}
  .chart-card h3 {{ margin: 0 0 10px; font-size: 0.95rem; }}
  .chart-wrap {{ position: relative; height: 240px; }}
  table {{ width: 100%; border-collapse: collapse; font-size: 0.85rem; background: var(--card); border: 1px solid var(--border); border-radius: 10px; overflow: hidden; }}
  th, td {{ padding: 8px 12px; text-align: left; border-bottom: 1px solid var(--border); }}
  th {{ background: var(--neutral-bg); color: var(--muted); font-weight: 600; }}
  tr:last-child td {{ border-bottom: none; }}
  .pill {{ font-size: 0.7rem; font-weight: 600; padding: 2px 8px; border-radius: 999px; }}
  .pill.good {{ background: var(--good-bg); color: var(--good); }}
  .pill.bad {{ background: var(--bad-bg); color: var(--bad); }}
  .empty {{ color: var(--muted); font-size: 0.85rem; padding: 8px 0; }}
</style>
</head>
<body>
<header>
  <div>
    <h1>{heading}</h1>
    <div class="meta">Fulfilment Plus &middot; generated {generated}</div>
  </div>
  <nav class="tabs">
    <a href="customer_dashboard.html" class="{cust_active}">Customer view</a>
    <a href="business_dashboard.html" class="{biz_active}">Business view</a>
  </nav>
</header>
<main>
{body}
</main>
<script>
const DATA = {data_json};
{script}
</script>
</body>
</html>
"""


def fmt_pct(x):
    return f"{x*100:.1f}%" if x is not None else "&ndash;"


def month_label(m):
    y, mo = m.split("-")
    return datetime.date(int(y), int(mo), 1).strftime("%b %Y")


def build_customer_dashboard(data):
    body = """
  <div class="controls">
    <label for="customer-select"><strong>Customer:</strong></label>
    <select id="customer-select" onchange="render()"></select>
  </div>
  <div id="cards" class="grid"></div>
  <div id="charts"></div>
  <div class="chart-section">
    <h2>Pick &amp; Pack Times</h2>
    <div id="pickpack"></div>
    <h3 style="margin:24px 0 10px;">Avg Time Per Order</h3>
    <div id="pickpack-per-order"></div>
  </div>
  <div class="chart-section">
    <h2>Action plan notes</h2>
    <div id="notes"></div>
  </div>
"""

    script = r"""
const monthLabel = (m) => {
  const [y, mo] = m.split('-');
  return new Date(y, mo - 1, 1).toLocaleString('en-NZ', { month: 'short', year: 'numeric' });
};

const asOfLabel = (iso) => {
  const d = new Date(iso);
  return d.toLocaleString('en-NZ', { day: 'numeric', month: 'short', year: 'numeric' });
};

// ── Circular KPI ring (this month's headline figure) ────────────────────
// pct: 0-100 or null (no data yet). frac: "num/den" string. met: whether
// pct clears the KPI's target - drives the ring colour (green/red).
function ringSvg(pct, frac, met) {
  const r = 49, c = 2 * Math.PI * r;
  if (pct === null) {
    return `<svg width="112" height="112" viewBox="0 0 112 112">
      <circle cx="56" cy="56" r="${r}" fill="none" class="ring-track ring-track-dashed"></circle>
      <text x="56" y="58" text-anchor="middle" class="ring-empty">No orders yet</text>
    </svg>`;
  }
  const filled = Math.max(0, Math.min(1, pct / 100));
  const offset = (c * (1 - filled)).toFixed(1);
  const colorClass = met ? 'ring-good' : 'ring-bad';
  return `<svg width="112" height="112" viewBox="0 0 112 112">
    <circle cx="56" cy="56" r="${r}" fill="none" class="ring-track"></circle>
    <circle cx="56" cy="56" r="${r}" fill="none" class="ring-progress ${colorClass}" stroke-dasharray="${c.toFixed(1)}" stroke-dashoffset="${offset}" transform="rotate(-90 56 56)"></circle>
    <text x="56" y="54" text-anchor="middle" class="ring-pct">${pct.toFixed(1)}%</text>
    <text x="56" y="70" text-anchor="middle" class="ring-frac">${frac}</text>
  </svg>`;
}

// ── Pick/pack time helpers ──────────────────────────────────────────────
// The time-tracking app's customerName doesn't always exactly match the
// KPI dashboard's customer names (e.g. "Jock Freight" vs "Jock Freight -
// SECA"). Known spelling mismatches that fuzzy matching can't catch go here.
const PICK_NAME_OVERRIDES = { "Coromandal Mountain": "Coromandel Mountains Company Limited" };

function normalizeName(s) {
  return (s || '').toLowerCase().replace(/[^a-z0-9]/g, '');
}

function findPickStats(customerName) {
  if (!DATA.pick_stats) return null;
  const norm = normalizeName(customerName);
  return DATA.pick_stats.find(p => {
    const aliased = PICK_NAME_OVERRIDES[p.customerName] || p.customerName;
    const pn = normalizeName(aliased);
    return pn === norm || norm.includes(pn) || pn.includes(norm);
  }) || null;
}

function fmtSeconds(totalSeconds) {
  if (totalSeconds === null || totalSeconds === undefined) return '–';
  const s = Math.round(totalSeconds);
  const h = Math.floor(s / 3600);
  const m = Math.floor((s % 3600) / 60);
  const sec = s % 60;
  const pad = (n) => String(n).padStart(2, '0');
  if (h > 0) return `${h}:${pad(m)}:${pad(sec)}`;
  if (m > 0) return `${m}:${pad(sec)}`;
  return `${sec}s`;
}

function renderPickPack(customer) {
  const el = document.getElementById('pickpack');
  if (!el) return;
  el.innerHTML = '';
  const stats = findPickStats(customer);
  if (!stats) {
    el.innerHTML = '<div class="empty">No pick/pack time data for this customer yet.</div>';
    return;
  }
  const grid = document.createElement('div');
  grid.className = 'grid';
  [
    { label: 'Previous month', data: stats.previousMonth },
    { label: 'This month to date', data: stats.currentMonthToDate },
  ].forEach(({ label, data }) => {
    ['picking', 'packing'].forEach(stage => {
      const s = data[stage];
      const stageLabel = stage === 'picking' ? 'Picking' : 'Packing';
      const card = document.createElement('div');
      card.className = 'card' + (!s ? ' nodata' : '');
      if (!s) {
        card.innerHTML = `<h3>${stageLabel}</h3>
          <div class="freq">${label}</div>
          <div class="empty">No orders recorded</div>`;
      } else {
        card.innerHTML = `<h3>${stageLabel}</h3>
          <div class="freq">${label}</div>
          <div class="figure">${fmtSeconds(s.avgSeconds)}</div>
          <div class="target">avg per time entry &middot; ${s.count} time entr${s.count === 1 ? 'y' : 'ies'} &middot; ${fmtSeconds(s.totalSeconds)} total</div>`;
      }
      grid.appendChild(card);
    });
  });
  el.appendChild(grid);
}

function mtdHtml(kpiDef, customer) {
  const mtd = DATA.mtd;
  if (!mtd) return '';
  const kpiData = mtd.kpis[kpiDef.kpi];
  if (!kpiData) return '';
  const rec = kpiData[customer];
  const label = `Month to date · ${monthLabel(mtd.month)} (as of ${asOfLabel(mtd.as_of)})`;
  if (!rec || rec.den === 0) {
    return `<div class="mtd"><div class="mtd-label">${label}</div><div class="mtd-na">No orders yet this month</div></div>`;
  }
  const pct = (rec.num / rec.den * 100).toFixed(1) + '%';
  return `<div class="mtd"><div class="mtd-label">${label}</div><div class="mtd-figure">${pct} <span class="mtd-frac">(${rec.num}/${rec.den})</span></div></div>`;
}

// ── Avg time per order (uses CartonCloud order counts already pulled in
// for the Order Cut Off KPI - "den" there is the total orders evaluated
// that month, used here as a proxy for orders shipped, so no extra
// CartonCloud API calls are needed). ─────────────────────────────────────
function orderCountFor(customer, mtdMode) {
  if (mtdMode) {
    const mtd = DATA.mtd;
    if (!mtd || !mtd.kpis['Order Cut Off']) return null;
    const rec = mtd.kpis['Order Cut Off'][customer];
    if (!rec || !rec.den) return null;
    return { den: rec.den, label: `${monthLabel(mtd.month)} MTD (as of ${asOfLabel(mtd.as_of)})` };
  }
  const recs = DATA.records
    .filter(r => r.customer === customer && r.kpi === 'Order Cut Off')
    .sort((a, b) => a.month.localeCompare(b.month));
  if (recs.length === 0) return null;
  const latest = recs[recs.length - 1];
  if (!latest.den) return null;
  return { den: latest.den, label: monthLabel(latest.month) };
}

function renderPickPackPerOrder(customer) {
  const el = document.getElementById('pickpack-per-order');
  if (!el) return;
  el.innerHTML = '';
  const stats = findPickStats(customer);
  if (!stats) {
    el.innerHTML = '<div class="empty">No pick/pack time data for this customer yet.</div>';
    return;
  }
  const grid = document.createElement('div');
  grid.className = 'grid';
  [
    { label: 'Previous month', data: stats.previousMonth, orderInfo: orderCountFor(customer, false) },
    { label: 'This month to date', data: stats.currentMonthToDate, orderInfo: orderCountFor(customer, true) },
  ].forEach(({ label, data, orderInfo }) => {
    ['picking', 'packing'].forEach(stage => {
      const s = data[stage];
      const stageLabel = stage === 'picking' ? 'Picking' : 'Packing';
      const card = document.createElement('div');
      if (!s || !orderInfo) {
        card.className = 'card nodata';
        const reason = !s ? 'No time entries recorded' : 'No CartonCloud order count for this period';
        card.innerHTML = `<h3>${stageLabel}</h3>
          <div class="freq">${label}</div>
          <div class="empty">${reason}</div>`;
      } else {
        const perOrder = s.totalSeconds / orderInfo.den;
        card.className = 'card';
        card.innerHTML = `<h3>${stageLabel}</h3>
          <div class="freq">${label} &middot; ${orderInfo.den} order${orderInfo.den === 1 ? '' : 's'} shipped (${orderInfo.label})</div>
          <div class="figure">${fmtSeconds(perOrder)}</div>
          <div class="target">avg per order &middot; ${fmtSeconds(s.totalSeconds)} total time entries</div>`;
      }
      grid.appendChild(card);
    });
  });
  el.appendChild(grid);
  const note = document.createElement('div');
  note.className = 'empty';
  note.style.marginTop = '8px';
  note.textContent = 'Order counts come from each customer\'s Order Cut Off KPI figures (already pulled from CartonCloud) as a proxy for orders shipped that month.';
  el.appendChild(note);
}

function populateSelect() {
  const sel = document.getElementById('customer-select');
  DATA.customers.forEach(c => {
    const opt = document.createElement('option');
    opt.value = c;
    opt.textContent = c;
    sel.appendChild(opt);
  });
}

let charts = [];

function render() {
  const customer = document.getElementById('customer-select').value;
  const cardsEl = document.getElementById('cards');
  const chartsEl = document.getElementById('charts');
  const notesEl = document.getElementById('notes');
  cardsEl.innerHTML = '';
  chartsEl.innerHTML = '';
  notesEl.innerHTML = '';
  charts.forEach(c => c.destroy());
  charts = [];

  const custRecords = DATA.records.filter(r => r.customer === customer);

  DATA.kpi_schedule.forEach(kpiDef => {
    const recs = custRecords.filter(r => r.kpi === kpiDef.kpi).sort((a, b) => a.month.localeCompare(b.month));
    const card = document.createElement('div');

    // MTD is now the headline figure on the card, with last month shown
    // smaller underneath - computed up front so both branches below can use it.
    const mtd = DATA.mtd;
    const mtdRec = (mtd && mtd.kpis[kpiDef.kpi]) ? mtd.kpis[kpiDef.kpi][customer] : null;
    const mtdPct = (mtdRec && mtdRec.den > 0) ? (mtdRec.num / mtdRec.den) : null;
    const mtdMet = mtdPct !== null && mtdPct >= kpiDef.target;
    const mtdLabel = mtd ? `MTD · ${monthLabel(mtd.month)} (as of ${asOfLabel(mtd.as_of)})` : '';

    card.className = 'card ring-card';

    if (recs.length === 0) {
      if (mtdPct === null) {
        card.innerHTML = `${ringSvg(null, null, false)}
          <h3>${kpiDef.kpi}</h3>
          <div class="ring-target-label">Target ${(kpiDef.target*100).toFixed(1)}%</div>
          <div class="ring-last-label">No data yet</div>`;
      } else {
        card.innerHTML = `${ringSvg(mtdPct * 100, `${mtdRec.num}/${mtdRec.den}`, mtdMet)}
          <h3>${kpiDef.kpi}</h3>
          <div class="ring-target-label">Target ${(kpiDef.target*100).toFixed(1)}%</div>
          <div class="ring-last-label">No closed-month data yet</div>`;
      }
      cardsEl.appendChild(card);
      return;
    }
    const latest = recs[recs.length - 1];
    // Primary (ring) figure is always month-to-date - if this customer/KPI
    // has no orders yet this month, the ring shows its own "no orders yet"
    // state rather than silently substituting last month's number.
    const primaryPctNum = mtdPct !== null ? mtdPct * 100 : null;
    const primaryFrac = mtdPct !== null ? `${mtdRec.num}/${mtdRec.den}` : null;
    const primaryMet = mtdMet;
    // Carrier breakdown for DIFOT on customer dashboard
    let custCarrierHtml = '';
    let custDifotMonth = null;
    if (kpiDef.kpi === 'Freight DIFOT' && DATA.difot_carriers) {
      const allDifotMonths = Object.keys(DATA.difot_carriers).sort();
      custDifotMonth = [...allDifotMonths].reverse().find(m => DATA.difot_carriers[m][customer]);
      const custCarriers = custDifotMonth && DATA.difot_carriers[custDifotMonth][customer];
      if (custCarriers) {
        const carriers = custCarriers;
        const rows = Object.entries(carriers)
          .map(([name, s]) => ({ name, pct: s.den ? s.num / s.den : null, num: s.num, den: s.den,
                                  avgDays: (s.days_count > 0) ? s.days_total / s.days_count : null }))
          .sort((a, b) => (b.pct ?? -1) - (a.pct ?? -1));
        const rowsHtml = rows.map(c => {
          const pctStr = c.pct !== null ? (c.pct * 100).toFixed(1) + '%' : '–';
          const ok = c.pct !== null && c.pct >= kpiDef.target;
          const daysStr = c.avgDays !== null ? c.avgDays.toFixed(1) + 'd avg' : '';
          return `<tr>
            <td style="padding:3px 8px 3px 0;font-size:0.8rem">${c.name}</td>
            <td style="padding:3px 0;font-size:0.8rem;font-weight:700;color:${ok ? '#16a34a' : '#dc2626'}">${pctStr}</td>
            <td style="padding:3px 0 3px 8px;font-size:0.75rem;color:var(--muted)">${c.num}/${c.den}${daysStr ? ' · ' + daysStr : ''}</td>
          </tr>`;
        }).join('');
        custCarrierHtml = `<table style="border-collapse:collapse;width:100%">${rowsHtml}</table>`;
      }
    }
    // Last-month line: always shown small under the target, as reference
    // context for the MTD headline figure in the ring above.
    const lastLineHtml = `Last month · ${latest.actual !== null ? (latest.actual*100).toFixed(1)+'%' : '–'}`;

    card.innerHTML = `${ringSvg(primaryPctNum, primaryFrac, primaryMet)}
      <h3>${kpiDef.kpi}</h3>
      <div class="ring-target-label">Target ${(kpiDef.target*100).toFixed(1)}%</div>
      <div class="ring-last-label">${lastLineHtml}</div>`;
    cardsEl.appendChild(card);

    // Carrier breakdown as a separate sibling card
    if (custCarrierHtml) {
      const carrierCard = document.createElement('div');
      carrierCard.className = 'card';
      carrierCard.innerHTML = `<h3 style="margin-bottom:8px">DIFOT by Carrier</h3>
        <div class="freq">${monthLabel(custDifotMonth)}</div>
        ${custCarrierHtml}`;
      cardsEl.appendChild(carrierCard);
    }

    if (latest.notes) {
      const noteDiv = document.createElement('div');
      noteDiv.className = 'chart-card';
      noteDiv.innerHTML = `<h3>${kpiDef.kpi} — ${monthLabel(latest.month)}</h3><div class="note" style="margin-top:0">${latest.notes}</div>`;
      notesEl.appendChild(noteDiv);
    }

    // Build trend labels + data, appending MTD point if available
    const custTrendLabels = recs.map(r => monthLabel(r.month));
    const custTrendData   = recs.map(r => r.actual !== null ? +(r.actual * 100).toFixed(1) : null);
    let custMtdPct = null;
    if (DATA.mtd && DATA.mtd.kpis[kpiDef.kpi] && DATA.mtd.kpis[kpiDef.kpi][customer]) {
      const m = DATA.mtd.kpis[kpiDef.kpi][customer];
      if (m.den > 0) custMtdPct = +(m.num / m.den * 100).toFixed(1);
    }
    if (custMtdPct !== null) {
      custTrendLabels.push(monthLabel(DATA.mtd.month) + ' MTD');
      custTrendData.push(custMtdPct);
    }
    const custLastCompIdx = custMtdPct !== null ? custTrendData.length - 2 : -1;

    const section = document.createElement('div');
    section.className = 'chart-card';
    section.innerHTML = `<h3>${kpiDef.kpi} trend</h3><div class="chart-wrap"><canvas></canvas></div>`;
    chartsEl.appendChild(section);
    const canvas = section.querySelector('canvas');
    const ctx = canvas.getContext('2d');
    const chart = new Chart(ctx, {
      type: 'line',
      data: {
        labels: custTrendLabels,
        datasets: [
          {
            label: custMtdPct !== null ? 'Actual (dashed = MTD)' : 'Actual',
            data: custTrendData,
            borderColor: '#1d4ed8',
            backgroundColor: 'rgba(29,78,216,0.1)',
            tension: 0.25,
            fill: true,
            pointRadius: custTrendData.map((_, i) => i === custTrendData.length - 1 && custMtdPct !== null ? 6 : 4),
            pointStyle: custTrendData.map((_, i) => i === custTrendData.length - 1 && custMtdPct !== null ? 'triangle' : 'circle'),
            segment: custLastCompIdx >= 0 ? {
              borderDash: (ctx) => ctx.p0DataIndex >= custLastCompIdx ? [5, 4] : undefined,
              borderColor: (ctx) => ctx.p0DataIndex >= custLastCompIdx ? '#93c5fd' : '#1d4ed8',
            } : undefined,
          },
          {
            label: 'Target',
            data: custTrendLabels.map(() => kpiDef.target * 100),
            borderColor: '#dc2626',
            borderDash: [6, 4],
            pointRadius: 0,
            fill: false,
          },
        ],
      },
      options: {
        responsive: true,
        maintainAspectRatio: false,
        scales: { y: { ticks: { callback: v => v + '%' } } },
      },
    });
    charts.push(chart);
  });

  if (notesEl.innerHTML === '') {
    notesEl.innerHTML = '<div class="empty">No outstanding action items for the latest month.</div>';
  }

  renderPickPack(customer);
  renderPickPackPerOrder(customer);
}

populateSelect();
render();
"""

    customers_with_data = sorted({r["customer"] for r in data["records"]})
    all_customers = data["customers"] if data["customers"] else customers_with_data
    payload = {
        "records": data["records"],
        "kpi_schedule": data["kpi_schedule"],
        "customers": all_customers,
        "mtd": data["mtd"],
        "difot_carriers": data["difot_carriers"],
        "pick_stats": data["pick_stats"],
    }
    return PAGE_SHELL.format(
        title="Customer KPI Dashboard - Fulfilment Plus",
        heading="Customer KPI Dashboard",
        generated=datetime.date.today().isoformat(),
        cust_active="active",
        biz_active="",
        body=body,
        data_json=json.dumps(payload),
        script=script,
    )


def build_business_dashboard(data):
    body = """
  <div class="chart-section">
    <h2>Overall performance</h2>
    <div id="summary" class="grid"></div>
  </div>
  <div class="chart-section">
    <h2>By customer (latest month)</h2>
    <div id="customer-charts"></div>
  </div>
  <div class="chart-section">
    <h2>Trend over time (overall)</h2>
    <div id="trend-charts"></div>
  </div>
  <div class="chart-section">
    <h2>Freight DIFOT &mdash; carrier performance by month</h2>
    <div id="carrier-trend"></div>
  </div>
  <div class="chart-section">
    <h2>Pick &amp; Pack Times &mdash; overview (previous month)</h2>
    <h3 style="margin:0 0 10px;">Avg Time Per Time Entry</h3>
    <div id="pickpack-summary" class="grid"></div>
    <div id="pickpack-chart"></div>
    <h3 style="margin:24px 0 10px;">Avg Time Per Order</h3>
    <div id="pickpack-per-order-summary" class="grid"></div>
    <div id="pickpack-per-order-chart"></div>
  </div>
"""

    script = r"""
const monthLabel = (m) => {
  const [y, mo] = m.split('-');
  return new Date(y, mo - 1, 1).toLocaleString('en-NZ', { month: 'short', year: 'numeric' });
};

const asOfLabel = (iso) => {
  const d = new Date(iso);
  return d.toLocaleString('en-NZ', { day: 'numeric', month: 'short', year: 'numeric' });
};

// ── Circular KPI ring (this month's headline figure) ────────────────────
// pct: 0-100 or null (no data yet). frac: "num/den" string. met: whether
// pct clears the KPI's target - drives the ring colour (green/red).
function ringSvg(pct, frac, met) {
  const r = 49, c = 2 * Math.PI * r;
  if (pct === null) {
    return `<svg width="112" height="112" viewBox="0 0 112 112">
      <circle cx="56" cy="56" r="${r}" fill="none" class="ring-track ring-track-dashed"></circle>
      <text x="56" y="58" text-anchor="middle" class="ring-empty">No orders yet</text>
    </svg>`;
  }
  const filled = Math.max(0, Math.min(1, pct / 100));
  const offset = (c * (1 - filled)).toFixed(1);
  const colorClass = met ? 'ring-good' : 'ring-bad';
  return `<svg width="112" height="112" viewBox="0 0 112 112">
    <circle cx="56" cy="56" r="${r}" fill="none" class="ring-track"></circle>
    <circle cx="56" cy="56" r="${r}" fill="none" class="ring-progress ${colorClass}" stroke-dasharray="${c.toFixed(1)}" stroke-dashoffset="${offset}" transform="rotate(-90 56 56)"></circle>
    <text x="56" y="54" text-anchor="middle" class="ring-pct">${pct.toFixed(1)}%</text>
    <text x="56" y="70" text-anchor="middle" class="ring-frac">${frac}</text>
  </svg>`;
}

// ── Pick/pack time helpers ──────────────────────────────────────────────
const PICK_NAME_OVERRIDES = { "Coromandal Mountain": "Coromandel Mountains Company Limited" };

function normalizeName(s) {
  return (s || '').toLowerCase().replace(/[^a-z0-9]/g, '');
}

function fmtSeconds(totalSeconds) {
  if (totalSeconds === null || totalSeconds === undefined) return '–';
  const s = Math.round(totalSeconds);
  const h = Math.floor(s / 3600);
  const m = Math.floor((s % 3600) / 60);
  const sec = s % 60;
  const pad = (n) => String(n).padStart(2, '0');
  if (h > 0) return `${h}:${pad(m)}:${pad(sec)}`;
  if (m > 0) return `${m}:${pad(sec)}`;
  return `${sec}s`;
}

function renderPickPackOverview() {
  const summaryEl = document.getElementById('pickpack-summary');
  const chartEl = document.getElementById('pickpack-chart');
  if (!summaryEl || !DATA.pick_stats) {
    if (summaryEl) summaryEl.innerHTML = '<div class="empty">No pick/pack time data yet.</div>';
    return;
  }

  ['picking', 'packing'].forEach(stage => {
    let totalSeconds = 0, totalCount = 0, customersWithData = 0;
    DATA.pick_stats.forEach(p => {
      const s = p.previousMonth[stage];
      if (s) { totalSeconds += s.totalSeconds; totalCount += s.count; customersWithData++; }
    });
    const stageLabel = stage === 'picking' ? 'Avg Picking Time' : 'Avg Packing Time';
    const card = document.createElement('div');
    card.className = 'card';
    if (totalCount === 0) {
      card.innerHTML = `<h3>${stageLabel}</h3><div class="freq">Previous month</div><div class="empty">No data yet</div>`;
    } else {
      const avg = totalSeconds / totalCount;
      card.innerHTML = `<h3>${stageLabel}</h3>
        <div class="freq">Previous month &middot; ${customersWithData} customer(s)</div>
        <div class="figure">${fmtSeconds(avg)}</div>
        <div class="target">${totalCount} time entries &middot; ${fmtSeconds(totalSeconds)} total</div>`;
    }
    summaryEl.appendChild(card);
  });

  const rows = DATA.pick_stats.map(p => {
    const pick = p.previousMonth.picking;
    const pack = p.previousMonth.packing;
    if (!pick && !pack) return null;
    const avgTotal = (pick ? pick.avgSeconds : 0) + (pack ? pack.avgSeconds : 0);
    return { name: PICK_NAME_OVERRIDES[p.customerName] || p.customerName, avgTotal, pick: pick ? pick.avgSeconds : 0, pack: pack ? pack.avgSeconds : 0 };
  }).filter(Boolean).sort((a, b) => b.avgTotal - a.avgTotal);

  if (rows.length > 0 && chartEl) {
    const section = document.createElement('div');
    section.className = 'chart-card';
    section.innerHTML = `<h3>Avg pick + pack time by customer &mdash; previous month</h3><div class="chart-wrap" style="height:${Math.max(180, rows.length * 34)}px"><canvas></canvas></div>`;
    chartEl.appendChild(section);
    const ctx = section.querySelector('canvas').getContext('2d');
    new Chart(ctx, {
      type: 'bar',
      data: {
        labels: rows.map(r => r.name),
        datasets: [
          { label: 'Picking', data: rows.map(r => r.pick), backgroundColor: '#1d4ed8' },
          { label: 'Packing', data: rows.map(r => r.pack), backgroundColor: '#7c3aed' },
        ],
      },
      options: {
        indexAxis: 'y',
        responsive: true,
        maintainAspectRatio: false,
        plugins: {
          legend: { position: 'bottom' },
          tooltip: {
            callbacks: {
              label: (ctx) => `${ctx.dataset.label}: ${fmtSeconds(ctx.parsed.x)}`,
            },
          },
        },
        scales: {
          x: { stacked: true, title: { display: true, text: 'time (h:mm:ss)' }, ticks: { callback: v => fmtSeconds(v) } },
          y: { stacked: true, ticks: { autoSkip: false } },
        },
      },
    });
  } else if (chartEl) {
    chartEl.innerHTML = '<p style="color:var(--muted)">No pick/pack time data yet.</p>';
  }
}

// ── Avg time per order (uses CartonCloud order counts already pulled in
// for the Order Cut Off KPI - "den" there is the total orders evaluated
// that month, used here as a proxy for orders shipped, so no extra
// CartonCloud API calls are needed). ─────────────────────────────────────
function orderCountForBiz(customer) {
  const recs = DATA.records
    .filter(r => r.customer === customer && r.kpi === 'Order Cut Off')
    .sort((a, b) => a.month.localeCompare(b.month));
  if (recs.length === 0) return null;
  const latest = recs[recs.length - 1];
  if (!latest.den) return null;
  return { den: latest.den, label: monthLabel(latest.month) };
}

function renderPickPackPerOrderOverview() {
  const summaryEl = document.getElementById('pickpack-per-order-summary');
  const chartEl = document.getElementById('pickpack-per-order-chart');
  if (!summaryEl || !DATA.pick_stats) {
    if (summaryEl) summaryEl.innerHTML = '<div class="empty">No pick/pack time data yet.</div>';
    return;
  }

  ['picking', 'packing'].forEach(stage => {
    let totalSeconds = 0, totalOrders = 0, customersWithData = 0;
    DATA.pick_stats.forEach(p => {
      const canonicalName = PICK_NAME_OVERRIDES[p.customerName] || p.customerName;
      const orderInfo = orderCountForBiz(canonicalName);
      const s = p.previousMonth[stage];
      if (s && orderInfo) { totalSeconds += s.totalSeconds; totalOrders += orderInfo.den; customersWithData++; }
    });
    const stageLabel = stage === 'picking' ? 'Avg Picking Time / Order' : 'Avg Packing Time / Order';
    const card = document.createElement('div');
    card.className = 'card';
    if (totalOrders === 0) {
      card.innerHTML = `<h3>${stageLabel}</h3><div class="freq">Previous month</div><div class="empty">No order-count data yet</div>`;
    } else {
      const avg = totalSeconds / totalOrders;
      card.innerHTML = `<h3>${stageLabel}</h3>
        <div class="freq">Previous month &middot; ${customersWithData} customer(s), ${totalOrders} orders</div>
        <div class="figure">${fmtSeconds(avg)}</div>
        <div class="target">${fmtSeconds(totalSeconds)} total time entries</div>`;
    }
    summaryEl.appendChild(card);
  });

  const rows = DATA.pick_stats.map(p => {
    const canonicalName = PICK_NAME_OVERRIDES[p.customerName] || p.customerName;
    const orderInfo = orderCountForBiz(canonicalName);
    const pick = p.previousMonth.picking;
    const pack = p.previousMonth.packing;
    if (!orderInfo || (!pick && !pack)) return null;
    const pickPerOrder = pick ? pick.totalSeconds / orderInfo.den : 0;
    const packPerOrder = pack ? pack.totalSeconds / orderInfo.den : 0;
    return {
      name: canonicalName,
      avgTotal: pickPerOrder + packPerOrder,
      pick: pickPerOrder,
      pack: packPerOrder,
      orderCount: orderInfo.den,
      orderMonth: orderInfo.label,
    };
  }).filter(Boolean).sort((a, b) => b.avgTotal - a.avgTotal);

  if (rows.length > 0 && chartEl) {
    const section = document.createElement('div');
    section.className = 'chart-card';
    section.innerHTML = `<h3>Avg pick + pack time per order by customer &mdash; previous month</h3><div class="chart-wrap" style="height:${Math.max(180, rows.length * 34)}px"><canvas></canvas></div>`;
    chartEl.appendChild(section);
    const ctx = section.querySelector('canvas').getContext('2d');
    new Chart(ctx, {
      type: 'bar',
      data: {
        labels: rows.map(r => r.name),
        datasets: [
          { label: 'Picking', data: rows.map(r => r.pick), backgroundColor: '#1d4ed8' },
          { label: 'Packing', data: rows.map(r => r.pack), backgroundColor: '#7c3aed' },
        ],
      },
      options: {
        indexAxis: 'y',
        responsive: true,
        maintainAspectRatio: false,
        plugins: {
          legend: { position: 'bottom' },
          tooltip: {
            callbacks: {
              label: (ctx) => `${ctx.dataset.label}: ${fmtSeconds(ctx.parsed.x)}`,
              afterLabel: (ctx) => `${rows[ctx.dataIndex].orderCount} orders (${rows[ctx.dataIndex].orderMonth})`,
            },
          },
        },
        scales: {
          x: { stacked: true, title: { display: true, text: 'time per order (h:mm:ss)' }, ticks: { callback: v => fmtSeconds(v) } },
          y: { stacked: true, ticks: { autoSkip: false } },
        },
      },
    });
  } else if (chartEl) {
    chartEl.innerHTML = '<p style="color:var(--muted)">No order-count data available yet.</p>';
  }

  const note = document.createElement('div');
  note.className = 'empty';
  note.style.marginTop = '8px';
  note.textContent = 'Order counts come from each customer\'s Order Cut Off KPI figures (already pulled from CartonCloud) as a proxy for orders shipped that month.';
  if (chartEl) chartEl.appendChild(note);
}

renderPickPackOverview();
renderPickPackPerOrderOverview();

function mtdHtml(kpiDef) {
  const mtd = DATA.mtd;
  if (!mtd) return '';
  const kpiData = mtd.kpis[kpiDef.kpi];
  if (!kpiData) return '';
  const label = `Month to date · ${monthLabel(mtd.month)} (as of ${asOfLabel(mtd.as_of)})`;
  const totalNum = Object.values(kpiData).reduce((s, r) => s + r.num, 0);
  const totalDen = Object.values(kpiData).reduce((s, r) => s + r.den, 0);
  if (totalDen === 0) {
    return `<div class="mtd"><div class="mtd-label">${label}</div><div class="mtd-na">No orders yet this month</div></div>`;
  }
  const pct = (totalNum / totalDen * 100).toFixed(1) + '%';
  return `<div class="mtd"><div class="mtd-label">${label}</div><div class="mtd-figure">${pct} <span class="mtd-frac">(${totalNum}/${totalDen})</span></div></div>`;
}

const months = [...new Set(DATA.records.map(r => r.month))].sort();
const latestMonth = months[months.length - 1];

const summaryEl = document.getElementById('summary');
const custChartsEl = document.getElementById('customer-charts');
const trendChartsEl = document.getElementById('trend-charts');

DATA.kpi_schedule.forEach(kpiDef => {
  const allRecs = DATA.records.filter(r => r.kpi === kpiDef.kpi);

  // MTD aggregate across all customers - computed up front, now the
  // headline figure on the card, with last month shown smaller underneath.
  const mtd = DATA.mtd;
  const mtdKpiData = (mtd && mtd.kpis[kpiDef.kpi]) ? mtd.kpis[kpiDef.kpi] : null;
  const mtdTotalNum = mtdKpiData ? Object.values(mtdKpiData).reduce((s, r) => s + r.num, 0) : 0;
  const mtdTotalDen = mtdKpiData ? Object.values(mtdKpiData).reduce((s, r) => s + r.den, 0) : 0;
  const mtdOverall = mtdTotalDen > 0 ? mtdTotalNum / mtdTotalDen : null;
  const mtdMet = mtdOverall !== null && mtdOverall >= kpiDef.target;
  const mtdLabel = mtd ? `MTD · ${monthLabel(mtd.month)} (as of ${asOfLabel(mtd.as_of)})` : '';
  const mtdCustCount = mtdKpiData ? Object.values(mtdKpiData).filter(r => r.den > 0).length : 0;
  const mtdMetCount = mtdKpiData
    ? Object.values(mtdKpiData).filter(r => r.den > 0 && (r.num / r.den) >= kpiDef.target).length
    : 0;

  if (allRecs.length === 0) {
    const card = document.createElement('div');
    card.className = 'card ring-card';
    if (mtdOverall === null) {
      card.innerHTML = `${ringSvg(null, null, false)}
        <h3>${kpiDef.kpi}</h3>
        <div class="ring-target-label">Target ${(kpiDef.target*100).toFixed(1)}%</div>
        <div class="ring-last-label">No data yet</div>`;
    } else {
      card.innerHTML = `${ringSvg(mtdOverall * 100, `${mtdTotalNum}/${mtdTotalDen}`, mtdMet)}
        <h3>${kpiDef.kpi}</h3>
        <div class="ring-target-label">Target ${(kpiDef.target*100).toFixed(1)}% · ${mtdMetCount}/${mtdCustCount} customers met</div>
        <div class="ring-last-label">No closed-month data yet</div>`;
    }
    summaryEl.appendChild(card);
    return;
  }

  const kpiMonths = [...new Set(allRecs.map(r => r.month))].sort();
  const kpiLatestMonth = kpiMonths[kpiMonths.length - 1];
  const latestRecs = allRecs.filter(r => r.month === kpiLatestMonth);
  const totalNum = latestRecs.reduce((s, r) => s + r.num, 0);
  const totalDen = latestRecs.reduce((s, r) => s + r.den, 0);
  const overall = totalDen ? totalNum / totalDen : null;
  const met = overall !== null && overall >= kpiDef.target;
  const metCount = latestRecs.filter(r => r.met).length;

  // Primary (ring) figure is always month-to-date - if no orders have been
  // evaluated yet this month, the ring shows its own "no orders yet" state
  // rather than silently substituting last month's number.
  const primaryPctNum = mtdOverall !== null ? mtdOverall * 100 : null;
  const primaryFrac = mtdOverall !== null ? `${mtdTotalNum}/${mtdTotalDen}` : null;
  const primaryCustCount = mtdCustCount;
  const primaryMetCount = mtdMetCount;
  const primaryMet = mtdMet;

  const card = document.createElement('div');
  card.className = 'card ring-card';

  // Carrier breakdown for DIFOT
  let carrierHtml = '';
  const difotCarrierMonth = DATA.difot_carriers
    ? Object.keys(DATA.difot_carriers).sort().pop()
    : null;
  if (kpiDef.kpi === 'Freight DIFOT' && difotCarrierMonth && DATA.difot_carriers[difotCarrierMonth] && DATA.difot_carriers[difotCarrierMonth]['_total']) {
    const carriers = DATA.difot_carriers[difotCarrierMonth]['_total'];
    const rows = Object.entries(carriers)
      .map(([name, s]) => ({ name, pct: s.den ? s.num / s.den : null, num: s.num, den: s.den,
                              avgDays: (s.days_count > 0) ? s.days_total / s.days_count : null }))
      .sort((a, b) => (b.pct ?? -1) - (a.pct ?? -1));
    const rowsHtml = rows.map(c => {
      const pctStr = c.pct !== null ? (c.pct * 100).toFixed(1) + '%' : '–';
      const ok = c.pct !== null && c.pct >= kpiDef.target;
      const daysStr = c.avgDays !== null ? c.avgDays.toFixed(1) + 'd avg' : '';
      return `<tr>
        <td style="padding:3px 8px 3px 0;font-size:0.8rem">${c.name}</td>
        <td style="padding:3px 0;font-size:0.8rem;font-weight:700;color:${ok ? '#16a34a' : '#dc2626'}">${pctStr}</td>
        <td style="padding:3px 0 3px 8px;font-size:0.75rem;color:var(--muted)">${c.num}/${c.den}${daysStr ? ' · ' + daysStr : ''}</td>
      </tr>`;
    }).join('');
    carrierHtml = `<table style="border-collapse:collapse;width:100%">${rowsHtml}</table>`;
  }

  // Last-month line: always shown small under the target, as reference
  // context for the MTD headline figure in the ring above.
  const bizLastLineHtml = `Last month · ${overall !== null ? (overall*100).toFixed(1)+'%' : '–'}`;

  card.innerHTML = `${ringSvg(primaryPctNum, primaryFrac, primaryMet)}
    <h3>${kpiDef.kpi}</h3>
    <div class="ring-target-label">Target ${(kpiDef.target*100).toFixed(1)}% · ${primaryMetCount}/${primaryCustCount} customers met</div>
    <div class="ring-last-label">${bizLastLineHtml}</div>`;
  summaryEl.appendChild(card);

  // Carrier breakdown as a separate sibling card
  if (carrierHtml) {
    const carrierCard = document.createElement('div');
    carrierCard.className = 'card';
    carrierCard.innerHTML = `<h3 style="margin-bottom:8px">DIFOT by Carrier</h3>
      <div class="freq">Overall · ${monthLabel(difotCarrierMonth)}</div>
      ${carrierHtml}`;
    summaryEl.appendChild(carrierCard);
  }

  if (latestRecs.length > 0) {
    const sorted = [...latestRecs].sort((a, b) => (a.actual ?? 0) - (b.actual ?? 0));
    const section = document.createElement('div');
    section.className = 'chart-card';
    section.innerHTML = `<h3>${kpiDef.kpi} — ${monthLabel(kpiLatestMonth)}</h3><div class="chart-wrap" style="height:${Math.max(180, sorted.length * 34)}px"><canvas></canvas></div>`;
    custChartsEl.appendChild(section);
    const ctx = section.querySelector('canvas').getContext('2d');
    new Chart(ctx, {
      type: 'bar',
      data: {
        labels: sorted.map(r => r.customer),
        datasets: [
          {
            label: 'Actual %',
            data: sorted.map(r => +(r.actual * 100).toFixed(1)),
            backgroundColor: sorted.map(r => r.met ? '#16a34a' : '#dc2626'),
          },
        ],
      },
      options: {
        indexAxis: 'y',
        responsive: true,
        maintainAspectRatio: false,
        plugins: { legend: { display: false } },
        scales: {
          x: { min: 0, max: 100, ticks: { callback: v => v + '%' } },
          y: { ticks: { autoSkip: false } },
        },
      },
    });
  }

  const trendData = months.map(m => {
    const recs = allRecs.filter(r => r.month === m);
    const num = recs.reduce((s, r) => s + r.num, 0);
    const den = recs.reduce((s, r) => s + r.den, 0);
    return den ? +(num / den * 100).toFixed(1) : null;
  });

  // Append MTD point if available
  const bizTrendLabels = months.map(monthLabel);
  const bizTrendData   = [...trendData];
  let bizMtdPct = null;
  if (DATA.mtd && DATA.mtd.kpis[kpiDef.kpi]) {
    const kpiMtd = DATA.mtd.kpis[kpiDef.kpi];
    const totalNum = Object.values(kpiMtd).reduce((s, r) => s + r.num, 0);
    const totalDen = Object.values(kpiMtd).reduce((s, r) => s + r.den, 0);
    if (totalDen > 0) bizMtdPct = +(totalNum / totalDen * 100).toFixed(1);
  }
  if (bizMtdPct !== null) {
    bizTrendLabels.push(monthLabel(DATA.mtd.month) + ' MTD');
    bizTrendData.push(bizMtdPct);
  }
  const bizLastCompIdx = bizMtdPct !== null ? bizTrendData.length - 2 : -1;

  const tSection = document.createElement('div');
  tSection.className = 'chart-card';
  tSection.innerHTML = `<h3>${kpiDef.kpi} — overall trend</h3><div class="chart-wrap"><canvas></canvas></div>`;
  trendChartsEl.appendChild(tSection);
  const tCtx = tSection.querySelector('canvas').getContext('2d');
  new Chart(tCtx, {
    type: 'line',
    data: {
      labels: bizTrendLabels,
      datasets: [
        {
          label: bizMtdPct !== null ? 'Overall actual (dashed = MTD)' : 'Overall actual',
          data: bizTrendData,
          borderColor: '#1d4ed8',
          backgroundColor: 'rgba(29,78,216,0.1)',
          tension: 0.25,
          fill: true,
          pointRadius: bizTrendData.map((_, i) => i === bizTrendData.length - 1 && bizMtdPct !== null ? 6 : 4),
          pointStyle: bizTrendData.map((_, i) => i === bizTrendData.length - 1 && bizMtdPct !== null ? 'triangle' : 'circle'),
          segment: bizLastCompIdx >= 0 ? {
            borderDash: (ctx) => ctx.p0DataIndex >= bizLastCompIdx ? [5, 4] : undefined,
            borderColor: (ctx) => ctx.p0DataIndex >= bizLastCompIdx ? '#93c5fd' : '#1d4ed8',
          } : undefined,
        },
        {
          label: 'Target',
          data: bizTrendLabels.map(() => kpiDef.target * 100),
          borderColor: '#dc2626',
          borderDash: [6, 4],
          pointRadius: 0,
          fill: false,
        },
      ],
    },
    options: {
      responsive: true,
      maintainAspectRatio: false,
      scales: { y: { ticks: { callback: v => v + '%' } } },
    },
  });
});

// ── Carrier trend chart ────────────────────────────────────────────────
const carrierTrendEl = document.getElementById('carrier-trend');
if (DATA.difot_carriers && carrierTrendEl) {
  const dcMonths = Object.keys(DATA.difot_carriers).sort();
  // collect all carrier names
  const carrierNames = [...new Set(
    dcMonths.flatMap(m => Object.keys(DATA.difot_carriers[m]['_total'] || {}))
  )].sort();

  const CARRIER_COLORS = [
    '#1d4ed8', '#16a34a', '#dc2626', '#d97706', '#7c3aed', '#0891b2'
  ];

  if (dcMonths.length > 0 && carrierNames.length > 0) {
    // find the relevant kpiDef for the target line
    const difotDef = DATA.kpi_schedule.find(k => k.kpi === 'Freight DIFOT');
    const targetPct = difotDef ? difotDef.target * 100 : null;

    const datasets = carrierNames.map((name, i) => ({
      label: name,
      data: dcMonths.map(m => {
        const s = (DATA.difot_carriers[m]['_total'] || {})[name];
        return (s && s.den) ? +(s.num / s.den * 100).toFixed(1) : null;
      }),
      borderColor: CARRIER_COLORS[i % CARRIER_COLORS.length],
      backgroundColor: CARRIER_COLORS[i % CARRIER_COLORS.length] + '22',
      tension: 0.25,
      fill: false,
      pointRadius: 5,
      pointHoverRadius: 7,
      spanGaps: true,
    }));

    if (targetPct !== null) {
      datasets.push({
        label: 'Target',
        data: dcMonths.map(() => targetPct),
        borderColor: '#94a3b8',
        borderDash: [6, 4],
        pointRadius: 0,
        fill: false,
        tension: 0,
      });
    }

    const section = document.createElement('div');
    section.className = 'chart-card';
    section.innerHTML = '<h3>DIFOT % by carrier</h3><div class="chart-wrap"><canvas></canvas></div>';
    carrierTrendEl.appendChild(section);
    new Chart(section.querySelector('canvas').getContext('2d'), {
      type: 'line',
      data: { labels: dcMonths.map(monthLabel), datasets },
      options: {
        responsive: true,
        maintainAspectRatio: false,
        plugins: {
          legend: { position: 'bottom' },
          tooltip: {
            callbacks: {
              afterLabel: (ctx) => {
                const name = ctx.dataset.label;
                const m = dcMonths[ctx.dataIndex];
                const s = (DATA.difot_carriers[m] || {})['_total']?.[name];
                if (!s) return '';
                const lines = [s.num + '/' + s.den + ' orders'];
                if (s.days_count > 0) lines.push((s.days_total / s.days_count).toFixed(1) + 'd avg delivery');
                return lines;
              }
            }
          }
        },
        scales: {
          y: { min: 0, max: 100, ticks: { callback: v => v + '%' } }
        },
      },
    });

    const tableSection = document.createElement('div');
    tableSection.className = 'chart-card';
    const thCells = '<tr><th style="text-align:left;padding:4px 12px 4px 0">Month</th>'
      + carrierNames.map(n => '<th style="padding:4px 8px;text-align:center">' + n + '</th>').join('') + '</tr>';
    const tbRows = dcMonths.map(m => {
      const cells = carrierNames.map(name => {
        const s = (DATA.difot_carriers[m]['_total'] || {})[name];
        if (!s || !s.den) return '<td style="padding:4px 8px;text-align:center;color:var(--muted)">-</td>';
        const pct = (s.num / s.den * 100).toFixed(1) + '%';
        const difotDef2 = DATA.kpi_schedule.find(k => k.kpi === 'Freight DIFOT');
        const ok = difotDef2 ? (s.num / s.den) >= difotDef2.target : true;
        return '<td style="padding:4px 8px;text-align:center;font-weight:600;color:' + (ok ? '#16a34a' : '#dc2626') + '">'
          + pct + '<br><span style="font-weight:400;font-size:0.75rem;color:var(--muted)">' + s.num + '/' + s.den + '</span></td>';
      }).join('');
      return '<tr><td style="padding:4px 12px 4px 0;font-size:0.85rem">' + monthLabel(m) + '</td>' + cells + '</tr>';
    }).join('');
    tableSection.innerHTML = '<h3 style="margin-bottom:8px">Carrier stats by month</h3>'
      + '<table style="border-collapse:collapse;width:100%;font-size:0.85rem">'
      + '<thead style="border-bottom:2px solid var(--border)">' + thCells + '</thead>'
      + '<tbody>' + tbRows + '</tbody></table>';
    carrierTrendEl.appendChild(tableSection);
  } else {
    carrierTrendEl.innerHTML = '<p style="color:var(--muted)">No carrier data yet.</p>';
  }
}
"""

    payload = {
        "records": data["records"],
        "kpi_schedule": data["kpi_schedule"],
        "mtd": data["mtd"],
        "difot_carriers": data["difot_carriers"],
        "pick_stats": data["pick_stats"],
    }
    return PAGE_SHELL.format(
        title="Business KPI Dashboard - Fulfilment Plus",
        heading="Business KPI Dashboard",
        generated=datetime.date.today().isoformat(),
        cust_active="",
        biz_active="active",
        body=body,
        data_json=json.dumps(payload),
        script=script,
    )


def main():
    data = load_data()
    # Exclude the current month from historical records — MTD badges cover it.
    current_month = datetime.date.today().strftime("%Y-%m")
    data["records"] = [r for r in data["records"] if r["month"] != current_month]
    with open("customer_dashboard.html", "w", encoding="utf-8") as f:
        f.write(build_customer_dashboard(data))
    with open("business_dashboard.html", "w", encoding="utf-8") as f:
        f.write(build_business_dashboard(data))
    print(f"Wrote customer_dashboard.html and business_dashboard.html "
          f"({len(data['records'])} KPI_Data rows, {len(data['customers'])} customers, "
          f"{len(data['kpi_schedule'])} KPIs)")


if __name__ == "__main__":
    main()
