"""
One-off edit of pipeline/kpi_dashboard.xlsx (Exceptions tab), run by the
exceptions-check workflow. Prints a before/after diff of the Exceptions tab
and row counts of every other sheet. Run from the pipeline/ directory.
"""

import copy
import datetime
import difflib
import sys
import zipfile

import openpyxl
from openpyxl.formatting.formatting import ConditionalFormattingList
from openpyxl.worksheet.cell_range import CellRange

WORKBOOK = "kpi_dashboard.xlsx"

ORDER_ROWS = [
    {
        "customer": "Salt Shark", "reference": "4938", "month": "2026-09",
        "reason": "Order not yet dispatched; confirmed by operations to be completed and dispatched "
                  "tomorrow (2 Oct 2026) - extension approved, not a genuine cutoff miss.",
    },
    {
        "customer": "Salt Shark", "reference": "4953", "month": "2026-09",
        "reason": "Order not yet dispatched; confirmed by operations to be completed and dispatched "
                  "tomorrow (2 Oct 2026) - extension approved, not a genuine cutoff miss.",
    },
    {
        "customer": "Flo & Frankie", "reference": "*", "month": None,
        "reason": "Customer agreement: Flo & Frankie Order Cut Off orders count as met going forward, "
                  "per agreement with the customer.",
    },
]
DATE_LOGGED = datetime.datetime(2026, 10, 1)

CONTENT_ROW = {
    "customer": "Vixxen", "match_field": "Delivery Company Name", "contains": "farmers",
    "rule": "Fixed (2026-10-13 20:00)",
    "from": datetime.datetime(2026, 10, 5), "to": datetime.datetime(2026, 10, 5),
}


def dump_sheet(ws):
    lines = []
    for row in ws.iter_rows():
        cells = []
        for c in row:
            if c.value is None:
                continue
            v = c.value
            if isinstance(v, datetime.datetime):
                v = v.strftime("%Y-%m-%d %H:%M") if (v.hour or v.minute) else v.strftime("%Y-%m-%d")
            cells.append(f"{c.coordinate}={v!s} [{c.number_format}]")
        if cells:
            lines.append(" | ".join(cells))
    return lines


def row_counts(wb):
    return {ws.title: (ws.max_row, ws.max_column) for ws in wb.worksheets}


def zip_parts(path):
    with zipfile.ZipFile(path) as z:
        names = z.namelist()
    return sorted(n for n in names if n.startswith(("xl/charts/", "xl/drawings/", "xl/media/", "xl/pivot")))


def find_row(ws, text, col=1, startswith=False):
    for r in range(1, ws.max_row + 1):
        v = ws.cell(row=r, column=col).value
        if isinstance(v, str) and (v.startswith(text) if startswith else v == text):
            return r
    return None


def copy_row_style(ws, src_row, dst_row, max_col):
    for col in range(1, max_col + 1):
        s = ws.cell(row=src_row, column=col)
        d = ws.cell(row=dst_row, column=col)
        if s.has_style:
            d.font = copy.copy(s.font)
            d.number_format = s.number_format
            d.alignment = copy.copy(s.alignment)
            d.border = copy.copy(s.border)
            d.fill = copy.copy(s.fill)
            d.protection = copy.copy(s.protection)


def shift_range_str(rng, at, n):
    """Shift a single A1 range string for an insert of n rows at row `at`."""
    cr = CellRange(rng)
    if cr.min_row >= at:
        cr.shift(row_shift=n)
    elif cr.max_row >= at:
        cr.expand(down=n)
    return cr.coord


def insert_rows_safely(ws, at, n):
    """insert_rows() that also moves merged cells, row heights, data
    validations and conditional formatting (openpyxl doesn't)."""
    spanning = [m for m in ws.merged_cells.ranges if m.min_row < at <= m.max_row]
    if spanning:
        sys.exit(f"ABORT: merged range(s) span the insert point: {spanning}")
    moving = [CellRange(m.coord) for m in ws.merged_cells.ranges if m.min_row >= at]
    for m in moving:
        ws.unmerge_cells(m.coord)

    heights = {r: d.height for r, d in ws.row_dimensions.items() if r >= at and d.height is not None}

    ws.insert_rows(at, n)

    for m in moving:
        m.shift(row_shift=n)
        ws.merge_cells(m.coord)

    for r in sorted(heights, reverse=True):
        ws.row_dimensions[r + n].height = heights[r]
    for r in range(at, at + n):
        ws.row_dimensions[r].height = ws.row_dimensions[at - 1].height

    for dv in ws.data_validations.dataValidation:
        dv.sqref = openpyxl.worksheet.cell_range.MultiCellRange(
            " ".join(shift_range_str(r.coord, at, n) for r in dv.sqref.ranges)
        )

    old_cf = ws.conditional_formatting
    new_cf = ConditionalFormattingList()
    for cf in old_cf:
        new_sqref = " ".join(shift_range_str(r.coord, at, n) for r in cf.sqref.ranges)
        for rule in cf.rules:
            new_cf.add(new_sqref, rule)
    ws.conditional_formatting = new_cf


def main():
    parts_before = zip_parts(WORKBOOK)
    wb = openpyxl.load_workbook(WORKBOOK)
    ws = wb["Exceptions"]
    before = dump_sheet(ws)
    counts_before = row_counts(wb)

    print("Merged ranges on Exceptions:", sorted(m.coord for m in ws.merged_cells.ranges))
    print("Data validations:", [(dv.type, dv.formula1, dv.sqref.coord if hasattr(dv.sqref, 'coord') else str(dv.sqref)) for dv in ws.data_validations.dataValidation])
    print("Conditional formats:", [str(cf.sqref) for cf in ws.conditional_formatting])
    print("Defined names:", list(wb.defined_names.keys()) if hasattr(wb.defined_names, "keys") else wb.defined_names)
    refs = []
    for other in wb.worksheets:
        for row in other.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("=") and "Exceptions" in c.value:
                    refs.append(f"{other.title}!{c.coordinate}: {c.value}")
    print("Formulas referencing Exceptions:", refs or "none")

    # ---- Idempotency guard ------------------------------------------------
    existing = set()
    for r in range(1, ws.max_row + 1):
        existing.add((ws.cell(r, 2).value, str(ws.cell(r, 3).value), ws.cell(r, 5).value))
    for o in ORDER_ROWS:
        if (o["customer"], o["reference"], "Order Cut Off") in existing:
            sys.exit(f"ABORT: order exception already present: {o['customer']} {o['reference']}")

    # ---- Part 1: Order-Level Exceptions -----------------------------------
    note_row = find_row(ws, "Action must be either", startswith=True)
    section2 = find_row(ws, "Order-Level Exceptions")
    if not note_row or not section2 or note_row < section2:
        sys.exit("ABORT: couldn't locate Order-Level Exceptions note row")
    print(f"\nOrder-Level note row is {note_row}; inserting {len(ORDER_ROWS)} rows above it")
    style_src = note_row - 1
    insert_rows_safely(ws, note_row, len(ORDER_ROWS))
    for i, o in enumerate(ORDER_ROWS):
        r = note_row + i
        copy_row_style(ws, style_src, r, 9)
        ws.cell(r, 1, DATE_LOGGED).number_format = "yyyy-mm-dd"
        ws.cell(r, 2, o["customer"])
        ws.cell(r, 3, o["reference"]).number_format = "@"
        mcell = ws.cell(r, 4, o["month"])
        mcell.number_format = "@"
        ws.cell(r, 5, "Order Cut Off")
        ws.cell(r, 6, "Count as met")
        ws.cell(r, 7, o["reason"])
        ws.cell(r, 8).value = None
        ws.cell(r, 9).value = None

    # ---- Part 2: Order Content-Based Overrides ----------------------------
    section3 = find_row(ws, "Order Content-Based Overrides")
    r = section3 + 2  # skip header row
    while ws.cell(r, 1).value:
        r += 1
    last = r - 1
    if any(ws.cell(r, c).value is not None for c in range(1, 10)):
        sys.exit(f"ABORT: row {r} after content overrides is not empty")
    blank_after = r + 1
    if ws.cell(blank_after, 1).value is not None:
        # No spare blank row - insert one so the note below stays intact.
        insert_rows_safely(ws, r, 1)
    print(f"Content-based override table ends at row {last}; writing new row at {r}")
    copy_row_style(ws, last, r, 9)
    ws.cell(r, 1, CONTENT_ROW["customer"])
    ws.cell(r, 2, CONTENT_ROW["match_field"])
    ws.cell(r, 3, CONTENT_ROW["contains"])
    ws.cell(r, 4, CONTENT_ROW["rule"])
    date_fmt = ws.cell(last, 5).number_format if ws.cell(last, 5).is_date else "yyyy-mm-dd"
    ws.cell(r, 5, CONTENT_ROW["from"]).number_format = date_fmt
    ws.cell(r, 6, CONTENT_ROW["to"]).number_format = date_fmt

    after = dump_sheet(ws)
    counts_after = row_counts(wb)
    wb.save(WORKBOOK)
    parts_after = zip_parts(WORKBOOK)

    print("\n===== Exceptions tab diff =====")
    for line in difflib.unified_diff(before, after, "before", "after", n=1, lineterm=""):
        print(line)

    print("\n===== Sheet sizes (rows x cols) before -> after =====")
    for name in counts_before:
        b, a = counts_before[name], counts_after[name]
        flag = "" if b == a else "   <-- CHANGED"
        print(f"  {name}: {b[0]}x{b[1]} -> {a[0]}x{a[1]}{flag}")

    print("\n===== Chart/drawing/media parts =====")
    print("  before:", parts_before)
    print("  after: ", parts_after)
    if parts_before != parts_after:
        print("  WARNING: chart/drawing parts differ after save")


if __name__ == "__main__":
    main()
